from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse
import json, os, re, tempfile
from email.parser import BytesParser
from email.policy import default
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.abspath(__file__))
PORT = int(os.environ.get('PORT', '8000'))

ALIASES = {
    'id': ['unit', 'unit id', 'unit number', 'unit no', 'unit #', 'id', 'code', 'apartment id'],
    'name': ['unit name', 'name', 'property name', 'apartment name', 'title'],
    'type': ['type', 'unit type', 'property type', 'category type'],
    'floor': ['floor', 'level', 'storey'],
    'area': ['area', 'area sqft', 'area (sqft)', 'area (sq ft)', 'sqft', 'sq ft', 'size'],
    'rate': ['rate', 'rate per sqft', 'rate per sq ft', 'rate / sq ft', 'price per sqft', 'price per sq ft'],
    'categoryCharge': ['category charge', 'category charges', 'category charge per sqft', 'category charges per sqft', 'extra charge', 'charges'],
    'status': ['status', 'availability', 'unit status']
}

def norm(v):
    return re.sub(r'[^a-z0-9]+', ' ', str(v or '').strip().lower()).strip()

def number(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or '').replace(',', '').replace('₨', '').replace('PKR', '').strip()
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    return float(m.group()) if m else 0

def match_header(h):
    n = norm(h)
    for key, names in ALIASES.items():
        if n in names or any(n == norm(x) for x in names):
            return key
    return None

def from_table(ws):
    rows = list(ws.iter_rows(values_only=True))
    best = None
    for ri, row in enumerate(rows[:20]):
        mapping = {}
        for ci, v in enumerate(row):
            k = match_header(v)
            if k and k not in mapping:
                mapping[k] = ci
        if len(mapping) >= 2 and ('id' in mapping or 'area' in mapping):
            best = (ri, mapping)
            break
    if not best:
        return []
    ri, mapping = best
    out = []
    for row in rows[ri+1:]:
        vals = [row[i] if i < len(row) else None for i in mapping.values()]
        if not any(v not in (None, '') for v in vals):
            continue
        def get(k, default=''):
            i = mapping.get(k)
            return row[i] if i is not None and i < len(row) else default
        uid = str(get('id','')).strip()
        if not uid:
            continue
        area = number(get('area',0)); rate = number(get('rate',0)); cat = number(get('categoryCharge',0))
        if area <= 0 and rate <= 0:
            continue
        out.append({
            'id': uid,
            'name': str(get('name') or f"{get('type') or 'Unit'} {uid}").strip(),
            'type': str(get('type') or 'Apartment').strip(),
            'floor': str(get('floor') or 'Ground').strip(),
            'area': area, 'rate': rate, 'categoryCharge': cat,
            'status': str(get('status') or 'Available').strip()
        })
    return out

def from_labeled_sheet(ws):
    data = {}
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v not in (None, '')]
        if len(vals) >= 2:
            data[norm(vals[0])] = vals[1]
    title = str(ws.title)
    m = re.search(r'\(([^)]+)\)', title)
    uid = m.group(1).strip() if m else ''
    area_m = re.search(r'(\d+(?:\.\d+)?)\s*(?:sq\s*ft|sqft)?', title, re.I)
    area = number(data.get('area sqft') or data.get('area sq ft') or (area_m.group(1) if area_m else 0))
    rate = number(data.get('rate per sqft') or data.get('rate per sq ft') or data.get('rate'))
    cat = number(data.get('category charges') or data.get('category charge'))
    if not uid or (area <= 0 and rate <= 0):
        return None
    return {
        'id': uid,
        'name': title,
        'type': str(data.get('unit type') or 'Apartment'),
        'floor': str(data.get('floor') or 'Ground'),
        'area': area, 'rate': rate, 'categoryCharge': cat,
        'status': str(data.get('status') or 'Available')
    }

def parse_workbook(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    units = []
    sheets = []
    for ws in wb.worksheets:
        found = from_table(ws)
        if found:
            units.extend(found); sheets.append(ws.title); continue
        one = from_labeled_sheet(ws)
        if one:
            units.append(one); sheets.append(ws.title)
    # de-duplicate by unit ID, keeping the latest occurrence
    dedup = {}
    for u in units:
        dedup[str(u['id'])] = u
    return {'units': list(dedup.values()), 'sheets': sheets}

class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=ROOT, **kwargs)
    def end_json(self, obj, code=200):
        raw = json.dumps(obj).encode('utf-8')
        self.send_response(code); self.send_header('Content-Type','application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(raw))); self.send_header('Access-Control-Allow-Origin','*'); self.end_headers(); self.wfile.write(raw)
    def do_POST(self):
        if urlparse(self.path).path != '/api/import-excel':
            return self.end_json({'error':'Not found'}, 404)
        try:
            length = int(self.headers.get('Content-Length','0') or 0)
            body = self.rfile.read(length)
            header_blob = ('Content-Type: ' + self.headers.get('Content-Type','') + '\r\nMIME-Version: 1.0\r\n\r\n').encode()
            msg = BytesParser(policy=default).parsebytes(header_blob + body)
            item = next((part for part in msg.iter_parts() if part.get_content_disposition() == 'form-data' and part.get_param('name', header='content-disposition') == 'file'), None)
            if item is None or not item.get_filename():
                return self.end_json({'error':'No file uploaded.'}, 400)
            filename = item.get_filename()
            suffix = os.path.splitext(filename)[1].lower()
            if suffix not in ('.xlsx', '.xlsm'):
                return self.end_json({'error':'Please upload an .xlsx or .xlsm Excel file.'}, 400)
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
                f.write(item.get_payload(decode=True)); path=f.name
            try:
                result = parse_workbook(path)
            finally:
                os.unlink(path)
            result['filename'] = filename
            result['count'] = len(result['units'])
            self.end_json(result)
        except Exception as e:
            self.end_json({'error': str(e)}, 500)

if __name__ == '__main__':
    print(f'EstateFlow running at http://localhost:{PORT}')
    print('Open that address in Chrome/Edge. Keep this window running while using Excel import.')
    ThreadingHTTPServer(('0.0.0.0', PORT), Handler).serve_forever()
